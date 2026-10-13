import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    # =========================================================================
    # SLIDE 1: EXECUTIVE DARK NAVY THEME
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)

    # Dark background
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(12, 22, 38) # Deep navy #0C1626
    bg1.line.fill.background()

    # Left profile container card
    left_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.6), Inches(3.6), Inches(6.3))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = RGBColor(19, 34, 58) # #13223A
    left_card.line.color.rgb = RGBColor(35, 60, 95)
    left_card.line.width = Pt(1.2)

    # Photo on left
    photo_path = 'public/hari-prakash-circle.png'
    if os.path.exists(photo_path):
        s1.shapes.add_picture(photo_path, Inches(1.3), Inches(0.9), Inches(2.2), Inches(2.2))

    # Name and Title inside left card
    tx_box = s1.shapes.add_textbox(Inches(0.75), Inches(3.2), Inches(3.3), Inches(3.5))
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p_name = tf.paragraphs[0]
    p_name.alignment = PP_ALIGN.CENTER
    r_name = p_name.add_run()
    r_name.text = "Dr G. Hari Prakash"
    r_name.font.name = "Georgia"
    r_name.font.size = Pt(19)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(255, 255, 255)

    p_role = tf.add_paragraph()
    p_role.alignment = PP_ALIGN.CENTER
    p_role.space_before = Pt(5)
    r_role = p_role.add_run()
    r_role.text = "Assistant Professor (Public Health)"
    r_role.font.name = "Calibri"
    r_role.font.size = Pt(12)
    r_role.font.bold = True
    r_role.font.color.rgb = RGBColor(64, 196, 196) # Bright teal

    p_inst = tf.add_paragraph()
    p_inst.alignment = PP_ALIGN.CENTER
    p_inst.space_before = Pt(3)
    r_inst = p_inst.add_run()
    r_inst.text = "Ramaiah School of Public Health\nMS Ramaiah Univ. of Applied Sciences, India"
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(10)
    r_inst.font.color.rgb = RGBColor(200, 215, 230)

    p_deg = tf.add_paragraph()
    p_deg.alignment = PP_ALIGN.CENTER
    p_deg.space_before = Pt(6)
    r_deg = p_deg.add_run()
    r_deg.text = "PhD (Public Health) · Pharm.D · MPH"
    r_deg.font.name = "Calibri"
    r_deg.font.size = Pt(10.5)
    r_deg.font.bold = True
    r_deg.font.color.rgb = RGBColor(230, 165, 95) # Amber/Gold

    # Divider in left card
    div = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(5.05), Inches(2.8), Pt(1))
    div.fill.solid()
    div.fill.fore_color.rgb = RGBColor(45, 75, 115)
    div.line.fill.background()

    # Contact Info on left
    tx_contact = s1.shapes.add_textbox(Inches(0.8), Inches(5.15), Inches(3.2), Inches(1.6))
    tfc = tx_contact.text_frame
    tfc.word_wrap = True
    tfc.margin_left = tfc.margin_right = tfc.margin_top = tfc.margin_bottom = 0

    contacts = [
        ("🌐  Website:", "www.drhari.co.in"),
        ("✉  Email:", "hariprakash607@gmail.com"),
        ("📱  Phone:", "+91 7306789341"),
        ("🆔  ORCID:", "0000-0002-9463-4084")
    ]
    for i, (label, val) in enumerate(contacts):
        p = tfc.paragraphs[0] if i == 0 else tfc.add_paragraph()
        p.space_before = Pt(3)
        r1 = p.add_run()
        r1.text = f"{label} "
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(160, 180, 205)
        r2 = p.add_run()
        r2.text = val
        r2.font.name = "Calibri"
        r2.font.size = Pt(9.5)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(235, 245, 255)

    # -------------------------------------------------------------------------
    # RIGHT SIDE: TOP METRICS ROW (4 Cards)
    # -------------------------------------------------------------------------
    metrics = [
        ("45+", "Publications", "Peer-reviewed papers (IF up to 56.7)"),
        ("10+", "Digital Health Tools", "Deployed platforms (Onco Care, Vichara)"),
        ("1", "Patent Published", "Medical Technology device patent"),
        ("ESMO / JSMO", "International Awards", "Travel Grants (Europe & Japan)")
    ]
    card_w = Inches(1.97)
    card_gap = Inches(0.18)
    start_x = Inches(4.5)
    top_y = Inches(0.6)
    card_h = Inches(1.2)

    for idx, (num, title, sub) in enumerate(metrics):
        cx = start_x + idx * (card_w + card_gap)
        c_shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, top_y, card_w, card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = RGBColor(20, 36, 62)
        c_shape.line.color.rgb = RGBColor(38, 70, 110)
        c_shape.line.width = Pt(1)

        ctf = c_shape.text_frame
        ctf.word_wrap = True
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ctf.margin_top = ctf.margin_bottom = ctf.margin_left = ctf.margin_right = Pt(4)

        p1 = ctf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        r_num = p1.add_run()
        r_num.text = num
        r_num.font.name = "Georgia"
        r_num.font.size = Pt(16 if len(num) > 4 else 18)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(230, 165, 95) # Amber

        p2 = ctf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(2)
        r_t = p2.add_run()
        r_t.text = title
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(9.5)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(240, 248, 255)

        p3 = ctf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        p3.space_before = Pt(1)
        r_s = p3.add_run()
        r_s.text = sub
        r_s.font.name = "Calibri"
        r_s.font.size = Pt(7.5)
        r_s.font.color.rgb = RGBColor(150, 175, 200)

    # -------------------------------------------------------------------------
    # RIGHT SIDE: MIDDLE TWO MAIN CARDS
    # -------------------------------------------------------------------------
    mid_y = Inches(2.0)
    mid_h = Inches(3.2)
    col_w = Inches(4.07)
    col2_x = start_x + col_w + Inches(0.28)

    # Left Column: Research Focus & Technical Expertise
    card_focus = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_x, mid_y, col_w, mid_h)
    card_focus.fill.solid()
    card_focus.fill.fore_color.rgb = RGBColor(18, 33, 56)
    card_focus.line.color.rgb = RGBColor(32, 60, 95)
    card_focus.line.width = Pt(1)

    tf_f = card_focus.text_frame
    tf_f.word_wrap = True
    tf_f.margin_top = Inches(0.18)
    tf_f.margin_left = Inches(0.22)
    tf_f.margin_right = Inches(0.22)
    tf_f.margin_bottom = Inches(0.18)

    p_fh = tf_f.paragraphs[0]
    r_fh = p_fh.add_run()
    r_fh.text = "RESEARCH & INNOVATION FOCUS"
    r_fh.font.name = "Calibri"
    r_fh.font.size = Pt(11)
    r_fh.font.bold = True
    r_fh.font.color.rgb = RGBColor(64, 196, 196)

    focus_points = [
        ("Digital Health & mHealth Interventions", "User-centred co-design, psychometric scales, symptom tracking & patient adherence platforms."),
        ("Oncology Care Delivery", "Scalable mHealth interventions to reduce cancer care disparities in low- & middle-income settings."),
        ("Health Technology Assessment (HTA)", "Cost-effectiveness, health economic modelling, and technology appraisal."),
        ("Implementation Science & AI for Health", "Machine-learning decision support, quasi-experimental & target trial emulation methods.")
    ]
    for title, desc in focus_points:
        p = tf_f.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        r1.text = f"▪  {title}: "
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(235, 245, 255)
        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "Calibri"
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(175, 195, 220)

    # Right Column: Career Experience & Roles
    card_exp = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col2_x, mid_y, col_w, mid_h)
    card_exp.fill.solid()
    card_exp.fill.fore_color.rgb = RGBColor(18, 33, 56)
    card_exp.line.color.rgb = RGBColor(32, 60, 95)
    card_exp.line.width = Pt(1)

    tf_e = card_exp.text_frame
    tf_e.word_wrap = True
    tf_e.margin_top = Inches(0.18)
    tf_e.margin_left = Inches(0.22)
    tf_e.margin_right = Inches(0.22)
    tf_e.margin_bottom = Inches(0.18)

    p_eh = tf_e.paragraphs[0]
    r_eh = p_eh.add_run()
    r_eh.text = "ACADEMIC & PROFESSIONAL ROLES"
    r_eh.font.name = "Calibri"
    r_eh.font.size = Pt(11)
    r_eh.font.bold = True
    r_eh.font.color.rgb = RGBColor(64, 196, 196)

    roles = [
        ("Assistant Professor (Public Health)", "Ramaiah School of Public Health, MSRUAS (2025–Present)", "Teaching MPH programme: Research Methodology, Biostatistics, Technology in Public Health, and NCDs."),
        ("Scientist-D / Senior Research Officer", "PGIMER, Chandigarh (Nov 2024–Jul 2025)", "Involved in developing online Health Technology Assessment (HTA) course."),
        ("Field Investigator", "World Health Organization (WHO) / JSS Medical College", "Multi-centre clinical trial on oxygen & respiratory support in COVID-19 in LMICs.")
    ]
    for r_title, r_inst, r_desc in roles:
        p1 = tf_e.add_paragraph()
        p1.space_before = Pt(6)
        r1 = p1.add_run()
        r1.text = f"▪  {r_title}"
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(235, 245, 255)

        p2 = tf_e.add_paragraph()
        r2 = p2.add_run()
        r2.text = f"   {r_inst}"
        r2.font.name = "Calibri"
        r2.font.size = Pt(8.5)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(230, 165, 95)

        p3 = tf_e.add_paragraph()
        r3 = p3.add_run()
        r3.text = f"   {r_desc}"
        r3.font.name = "Calibri"
        r3.font.size = Pt(8.5)
        r3.font.color.rgb = RGBColor(175, 195, 220)

    # -------------------------------------------------------------------------
    # RIGHT SIDE: BOTTOM ROW - EDITORIAL & GLOBAL RECOGNITION
    # -------------------------------------------------------------------------
    bot_y = Inches(5.38)
    bot_h = Inches(1.52)
    total_r_w = col_w * 2 + Inches(0.28)

    card_bot = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_x, bot_y, total_r_w, bot_h)
    card_bot.fill.solid()
    card_bot.fill.fore_color.rgb = RGBColor(18, 33, 56)
    card_bot.line.color.rgb = RGBColor(32, 60, 95)
    card_bot.line.width = Pt(1)

    tf_b = card_bot.text_frame
    tf_b.word_wrap = True
    tf_b.margin_top = Inches(0.14)
    tf_b.margin_left = Inches(0.22)
    tf_b.margin_right = Inches(0.22)
    tf_b.margin_bottom = Inches(0.12)

    p_bh = tf_b.paragraphs[0]
    r_bh = p_bh.add_run()
    r_bh.text = "EDITORIAL APPOINTMENTS & INTERNATIONAL AFFILIATIONS"
    r_bh.font.name = "Calibri"
    r_bh.font.size = Pt(10.5)
    r_bh.font.bold = True
    r_bh.font.color.rgb = RGBColor(64, 196, 196)

    p_b1 = tf_b.add_paragraph()
    p_b1.space_before = Pt(4)
    r_b1 = p_b1.add_run()
    r_b1.text = "• Associate Editor: "
    r_b1.font.name = "Calibri"
    r_b1.font.size = Pt(9.2)
    r_b1.font.bold = True
    r_b1.font.color.rgb = RGBColor(235, 245, 255)
    r_b2 = p_b1.add_run()
    r_b2.text = "Archives of Public Health (BMC / Springer Nature)   |   Clinical Epidemiology and Global Health (Elsevier)"
    r_b2.font.name = "Calibri"
    r_b2.font.size = Pt(9.2)
    r_b2.font.color.rgb = RGBColor(185, 205, 230)

    p_b2 = tf_b.add_paragraph()
    p_b2.space_before = Pt(3)
    r_b3 = p_b2.add_run()
    r_b3.text = "• Professional Memberships: "
    r_b3.font.name = "Calibri"
    r_b3.font.size = Pt(9.2)
    r_b3.font.bold = True
    r_b3.font.color.rgb = RGBColor(235, 245, 255)
    r_b4 = p_b2.add_run()
    r_b4.text = "European Society of Medical Oncology (ESMO)   |   ESMO & JSMO Travel Fellowship Awardee"
    r_b4.font.name = "Calibri"
    r_b4.font.size = Pt(9.2)
    r_b4.font.color.rgb = RGBColor(185, 205, 230)
    output_path = 'Dr_G_Hari_Prakash_Executive_Profile.pptx'
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == '__main__':
    create_deck()
